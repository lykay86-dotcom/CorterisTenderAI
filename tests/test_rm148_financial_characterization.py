"""RM-148 baseline characterization; assertions intentionally describe v2 defects."""

from __future__ import annotations

from decimal import Decimal
import json
from zipfile import ZipFile

from openpyxl import load_workbook
from PySide6.QtWidgets import QApplication

from app.core.workflow_database_health import (
    WorkflowDatabaseHealthService,
    WorkflowDatabaseHealthStatus,
)
from app.reporting.workflow_excel import WorkflowExcelExporter
from app.repositories.business_metrics import (
    BusinessMetricsRepository,
    BusinessRecordKind,
    BusinessStatus,
)
from app.ui.business_workflow.dialogs import BusinessRecordDialog
from app.ui.business_workflow.model import WorkflowTableModel
from app.ui.viewmodels.dashboard_viewmodel import DashboardKpiUnit, _format_raw_value


def _app() -> QApplication:
    return QApplication.instance() or QApplication([])


def test_v2_repository_round_trips_decimal_through_json_float(tmp_path) -> None:
    """Expected-red will replace the float types and ordinary JSON-number lexemes."""
    path = tmp_path / "business_workflow.json"
    repository = BusinessMetricsRepository(path)

    record = repository.save_record(
        kind=BusinessRecordKind.PROPOSAL,
        tender_id="AUDIT",
        title="Audit",
        status=BusinessStatus.READY,
        total=Decimal("0.10"),
        profit=Decimal("0.01"),
        margin_percent=Decimal("10.00"),
    )
    payload = json.loads(path.read_text(encoding="utf-8"))

    assert repository.SCHEMA_VERSION == 2
    assert isinstance(record.total, float)
    assert isinstance(record.profit, float)
    assert isinstance(record.margin_percent, float)
    assert payload["records"][0]["total"] == 0.1
    assert '"total": 0.1' in path.read_text(encoding="utf-8")


def test_v2_invalid_number_silently_becomes_zero() -> None:
    """Expected-red will require a typed INVALID state instead of zero."""
    assert BusinessMetricsRepository._number("not-a-number") == 0.0


def test_baseline_ui_financial_precision_is_inconsistent() -> None:
    """Expected-red will route all projections through one formatter."""
    _app()
    money_spin = BusinessRecordDialog._money_spin()

    assert money_spin.decimals() == 2
    assert WorkflowTableModel._money(Decimal("1.50")) == "2 ₽"
    assert _format_raw_value(Decimal("1.50"), DashboardKpiUnit.RUB) == "2 ₽"
    assert f"{Decimal('1.25'):.1f}%" == "1.2%"


def test_baseline_health_accepts_nonfinite_float_text(tmp_path) -> None:
    """Expected-red will reject non-finite persisted values through Decimal validation."""
    repository = BusinessMetricsRepository(tmp_path / "business_workflow.json")
    repository.path.write_text(
        json.dumps(
            {
                "schema_version": 2,
                "records": [
                    {
                        "id": "legacy",
                        "kind": "proposal",
                        "tender_id": "T-1",
                        "title": "Legacy",
                        "status": "ready",
                        "total": "NaN",
                        "profit": 0,
                        "margin_percent": 0,
                    }
                ],
                "events": [],
            }
        ),
        encoding="utf-8",
    )

    report = WorkflowDatabaseHealthService().inspect(repository)

    assert report.status is WorkflowDatabaseHealthStatus.HEALTHY


def test_baseline_xlsx_uses_only_binary_numeric_cell(tmp_path) -> None:
    """Expected-red will add an exact fixed-point metadata representation."""
    repository = BusinessMetricsRepository(tmp_path / "source.json")
    record = repository.save_record(
        kind=BusinessRecordKind.PROPOSAL,
        tender_id="AUDIT",
        title="Audit",
        status=BusinessStatus.READY,
        total=Decimal("0.10"),
        profit=Decimal("0.01"),
        margin_percent=Decimal("10.00"),
    )
    target = tmp_path / "workflow.xlsx"
    WorkflowExcelExporter().export(target, records=[record])

    workbook = load_workbook(target)
    registry = workbook.worksheets[1]
    with ZipFile(target) as archive:
        xml = archive.read("xl/worksheets/sheet2.xml").decode("utf-8")

    assert registry["F2"].value == 0.1
    assert isinstance(registry["F2"].value, float)
    assert '<c r="F2" s="12" t="n"><v>0.1</v></c>' in xml
    assert "FinancialExact" not in workbook.sheetnames
